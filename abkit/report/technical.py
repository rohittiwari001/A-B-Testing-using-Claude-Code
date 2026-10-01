"""Technical report (technical_report.docx): the data scientist's companion to summary.docx.

Where summary.docx explains the result to stakeholders, this document explains *how* it was obtained, so a
data scientist can audit, challenge or reproduce every step:

1. Technical summary        verdict, decision rule, primary estimate, caveats
2. Problem framing          formal hypotheses, design, metrics, statistical settings and overrides
3. Data and validation      lineage (source, copy, SHA-256), every check, split test
4. Methods                  per plan step: why, estimator / formula, inference, assumptions checked, corrections, script
5. Results                  all estimates with CIs, Bayesian outputs, diagnostic tests, charts
6. Robustness               the same metric across methods, distance to decision boundaries, diagnostics, caveats
7. Decision log             intake Q&A, user decisions at gates, plan changes, state history, review rounds
8. Reproducibility          environment, seeds, code index, rerun commands, artefacts, run log, promotion candidates

Verbatim records (intake answers, review findings) use Word's "Quote" style; ``abkit.tracecheck`` skips them
because they may quote numbers that were later superseded. Everything else is generated from results.json.

Example
-------
>>> from abkit.report.technical import build_technical_report
>>> build_technical_report("runs/2026-10-01_marketing-ads")       # doctest: +SKIP
"""

from __future__ import annotations

import hashlib
import json
import platform
import re
import sys
from importlib import metadata
from pathlib import Path
from typing import Any

from docx.shared import Pt

from abkit import results as R
from abkit.deck.builder import _validation_rows, intake_sections, run_context
from abkit.report.builder import Report, _first_doc_line, _figures

# --------------------------------------------------------------------------- method reference

METHOD_DOCS: dict[str, dict[str, Any]] = {
    "two_proportion_ztest": {
        "name": "Two-proportion z-test",
        "what": "Compares conversion rates between two independent groups of randomisation units.",
        "formula": ["p_hat = x / n per group;  diff = p_T - p_C",
                    "z = diff / sqrt( p_pool (1 - p_pool) (1/n_C + 1/n_T) ),  p_pool = (x_C + x_T) / (n_C + n_T)",
                    "CI(diff) = diff +/- z_crit * sqrt( p_C(1-p_C)/n_C + p_T(1-p_T)/n_T )   (unpooled)",
                    "rel = p_T/p_C - 1;  Var(rel) ~ se_T^2/p_C^2 + p_T^2 se_C^2/p_C^4   (delta method)"],
        "inference": "Pooled SE for the test, unpooled SE for the CI; one-sided tests use a (1 - 2 alpha) two-sided CI.",
        "assumptions": "Independent units (one row per randomisation unit); at least 10 successes and failures per group.",
        "reference": "abkit.frequentist.two_proportion_ztest; ab-methods/frequentist_tests.md"},
    "welch_ttest": {
        "name": "Welch's t-test",
        "what": "Compares means of a continuous or count metric without assuming equal variances.",
        "formula": ["diff = mean_T - mean_C;  se = sqrt( s_T^2/n_T + s_C^2/n_C )",
                    "t = diff / se;  df = se^4 / ( (s_T^2/n_T)^2/(n_T-1) + (s_C^2/n_C)^2/(n_C-1) )   (Welch-Satterthwaite)",
                    "CI(diff) = diff +/- t_crit(df) * se;  relative CI by the delta method"],
        "inference": "Large-sample t inference on the difference in means.",
        "assumptions": "Independent units; approximately normal sampling distribution of the mean (CLT; skew checked).",
        "reference": "abkit.frequentist.welch_ttest; ab-methods/frequentist_tests.md"},
    "bootstrap": {
        "name": "Percentile bootstrap",
        "what": "Resamples units with replacement within each group to get the sampling distribution of the difference.",
        "formula": ["for b = 1..B: diff_b = stat(resample_T) - stat(resample_C)",
                    "CI = [quantile(diff_b, alpha/2), quantile(diff_b, 1 - alpha/2)];  p = 2 min(P(diff_b <= 0), P(diff_b >= 0))"],
        "inference": "Non-parametric; used as a robustness check for heavy-tailed metrics.",
        "assumptions": "Independent units; the sample represents the population.",
        "reference": "abkit.frequentist.bootstrap_diff"},
    "mann_whitney": {
        "name": "Mann-Whitney U test", "what": "Rank test for a shift in the distribution (not the mean).",
        "formula": ["U = rank-sum statistic;  P(T > C) + 0.5 P(T = C) = U / (n_T n_C)"],
        "inference": "Exact / normal-approximation p-value from scipy.", "assumptions": "Independent units; ordinal outcome.",
        "reference": "abkit.frequentist.mann_whitney"},
    "chi_square_test": {
        "name": "Chi-square test of independence", "what": "Tests association between variant and a categorical outcome.",
        "formula": ["chi2 = sum (O - E)^2 / E;  Cramer's V = sqrt( chi2 / (N (min(r, c) - 1)) )"],
        "inference": "Asymptotic chi-square distribution.", "assumptions": "Expected count >= 5 in each cell.",
        "reference": "abkit.frequentist.chi_square_test"},
    "noninferiority_test": {
        "name": "Non-inferiority test (guardrail)",
        "what": "Tests that the treatment is not worse than control by more than a pre-agreed relative margin m.",
        "formula": ["H0: diff <= -m * mean_C  (no_decrease)   vs   H1: diff > -m * mean_C",
                    "z = (diff + m * mean_C) / se;  one-sided p = 1 - Phi(z);  CI at 1 - 2 alpha",
                    "status: pass if p < alpha; breach if the whole CI is beyond the margin; else inconclusive"],
        "inference": "One-sided normal test.", "assumptions": "Normal approximation for the difference.",
        "reference": "abkit.frequentist.noninferiority_test; ab-methods/guardrails.md"},
    "beta_binomial": {
        "name": "Bayesian beta-binomial model",
        "what": "Posterior distribution of each group's conversion rate; probability that treatment is better.",
        "formula": ["prior p ~ Beta(a0, b0);  posterior p ~ Beta(a0 + x, b0 + n - x)  per group",
                    "P(beat) = P(p_T > p_C) by Monte Carlo draws;  expected loss = E[max(p_C - p_T, 0)]",
                    "credible intervals = posterior quantiles of p_T - p_C and p_T/p_C - 1"],
        "inference": "Monte Carlo with a fixed seed; decision: P(beat) >= threshold and negligible expected loss.",
        "assumptions": "Independent units; prior fixed before seeing results (weakly informative by default).",
        "reference": "abkit.bayesian.beta_binomial; ab-methods/bayesian.md"},
    "normal_model": {
        "name": "Bayesian normal model", "what": "Posterior of the difference in means with a N(0, tau^2) prior.",
        "formula": ["d_hat ~ N(d, se^2);  prior d ~ N(0, tau^2), tau = 10% of the control mean",
                    "posterior mean = w d_hat, sd = sqrt(w) se,  w = tau^2 / (tau^2 + se^2)"],
        "inference": "Monte Carlo draws for P(beat), expected loss, credible intervals.", "assumptions": "Large-sample normal likelihood.",
        "reference": "abkit.bayesian.normal_model"},
    "cuped": {
        "name": "CUPED (variance reduction)",
        "what": "Removes the part of the metric explained by a pre-period covariate X before comparing groups.",
        "formula": ["theta = cov(Y, X) / var(X)  (pooled);  Y_adj = Y - theta (X - mean(X))",
                    "Welch test on Y_adj;  variance reduction ~ corr(X, Y)^2"],
        "inference": "Unbiased because X is measured before assignment; relative lift against the unadjusted control mean.",
        "assumptions": "Covariate pre-assignment and balanced across groups.",
        "reference": "abkit.variance_reduction.cuped; ab-methods/cuped.md"},
    "regression_adjustment": {
        "name": "Regression adjustment (Lin 2013)", "what": "OLS of Y on treatment, centred covariates and their interactions.",
        "formula": ["Y = b0 + b1 T + b2' Xc + b3' (T x Xc) + e;  effect = b1;  HC2 robust SEs"],
        "inference": "Robust t inference on b1.", "assumptions": "Covariates measured pre-assignment.",
        "reference": "abkit.variance_reduction.regression_adjustment"},
    "delta_method_ratio": {
        "name": "Delta method for a ratio metric", "what": "Ratio of sums with units aggregated to the randomisation unit.",
        "formula": ["R = mean(X)/mean(Y) over units;  Var(R) = [var X/mu_Y^2 - 2 mu_X cov(X,Y)/mu_Y^3 + mu_X^2 var Y/mu_Y^4] / n"],
        "inference": "Normal approximation; design effect vs the naive row-level variance reported.",
        "assumptions": "Randomisation units independent; enough units per arm.", "reference": "abkit.ratio_metrics.delta_method_ratio"},
    "cluster_robust_test": {
        "name": "Cluster-robust regression", "what": "Row-level OLS / WLS with SEs clustered by randomisation unit.",
        "formula": ["Y = b0 + b1 T + e;  CR1 cluster-robust covariance"], "inference": "Wald test on b1.",
        "assumptions": "Many clusters.", "reference": "abkit.ratio_metrics.cluster_robust_test"},
    "interaction_test": {
        "name": "Treatment x segment interaction test",
        "what": "Joint Wald test that the absolute treatment effect is equal in every segment.",
        "formula": ["Y = b0 + b1 T + sum_s g_s S_s + sum_s d_s (T x S_s) + e  (OLS, HC1);  H0: all d_s = 0"],
        "inference": "Wald chi-square; the stored 'estimate' is the test statistic, not an effect.",
        "assumptions": "Linear probability model for binary metrics.", "reference": "abkit.segments.interaction_test"},
    "simpsons_check": {
        "name": "Simpson's paradox check", "what": "Compares the pooled difference with a segment-standardised difference.",
        "formula": ["standardised = sum_s w_s (mean_T,s - mean_C,s),  w_s = overall segment share;  flag if signs differ"],
        "inference": "Descriptive diagnostic.", "assumptions": "None.", "reference": "abkit.segments.simpsons_check"},
    "srm_check": {
        "name": "Sample ratio mismatch test", "what": "Chi-square goodness of fit of unit counts against the intended split.",
        "formula": ["chi2 = sum_v (n_v - N s_v)^2 / (N s_v),  s_v = intended share;  SRM if p < threshold (default 0.001)"],
        "inference": "Asymptotic chi-square.", "assumptions": "Counts of randomisation units.", "reference": "abkit.validation.srm_check"},
    "trend_test": {
        "name": "Novelty / primacy trend test", "what": "Inverse-variance weighted regression of per-period lift on period index.",
        "formula": ["lift_t = a + b t + e,  weights 1/se_t^2;  decaying (novelty) if b is significant and opposite in sign to the lift"],
        "inference": "WLS t-test on b.", "assumptions": "Enough units per period.", "reference": "abkit.time_effects.trend_test"},
    "diff_in_diff": {
        "name": "Difference-in-differences", "what": "Effect on treated units relative to their counterfactual trend.",
        "formula": ["Y_it = a_i + l_t + b (treated_i x post_t) + e_it  (two-way fixed effects), SEs clustered by unit"],
        "inference": "Cluster-robust Wald test on b.", "assumptions": "Parallel trends; no spillover or anticipation.",
        "reference": "abkit.quasi.diff_in_diff"},
    "synthetic_control": {
        "name": "Synthetic control", "what": "Convex donor weights fitted to the pre-period; effect = post-period gap.",
        "formula": ["min_w sum_pre (y_treated - Y_donors w)^2  s.t. w >= 0, sum w = 1;  placebo p from RMSPE ratios"],
        "inference": "Placebo-in-space permutation p-value.", "assumptions": "Good pre-period fit.", "reference": "abkit.quasi.synthetic_control"},
    "interrupted_time_series": {
        "name": "Interrupted time series", "what": "Segmented regression of the outcome on time with a level and slope change.",
        "formula": ["y_t = b0 + b1 t + b2 post_t + b3 (t - t0) post_t + e_t;  Newey-West SEs"],
        "inference": "HAC t-tests.", "assumptions": "No co-occurring change.", "reference": "abkit.quasi.interrupted_time_series"},
    "msprt": {
        "name": "Mixture SPRT (always-valid inference)", "what": "Sequential test valid under continuous monitoring.",
        "formula": ["LR_n = sqrt(V/(V+tau^2)) exp( d^2 tau^2 / (2V(V+tau^2)) );  p_n = min(p_{n-1}, 1/LR_n)"],
        "inference": "Always-valid p-values and confidence sequences.", "assumptions": "tau set in advance.",
        "reference": "abkit.sequential.msprt_path"},
    "build_summary": {
        "name": "Decision rule and summary", "what": "Turns the primary result and guardrail statuses into a verdict.",
        "formula": ["significantly worse -> don't ship;  guardrail breach -> don't ship (iterate if the primary won)",
                    "significantly better and rel. lift >= threshold -> ship;  significant but < threshold -> iterate",
                    "not significant: CI excludes the threshold -> iterate;  otherwise -> extend (inconclusive)",
                    "Bayesian primary: P(beat) >= threshold counts as significantly better"],
        "inference": "Deterministic rule applied to stored results.", "assumptions": "A practical threshold agreed at intake.",
        "reference": "abkit.decision.recommend / build_summary; ab-methods/interpretation_and_decisions.md"},
    "full_profile": {
        "name": "Data validation profile", "what": "Schema, dtypes, missing values, duplicates, contamination, variants, SRM, "
                                                  "dates, outliers, impossible values, pre-period data, covariate balance.",
        "formula": [], "inference": "pass / warn / block per check.", "assumptions": "None.",
        "reference": "abkit.validation.full_profile"},
}


def method_doc(method: str) -> dict[str, Any] | None:
    key = (method or "").split(".")[-1]
    if key.startswith("bootstrap"):
        key = "bootstrap"
    return METHOD_DOCS.get(key)


# --------------------------------------------------------------------------- small helpers


def _code(rep: Report, lines: list[str]) -> None:
    for ln in lines:
        p = rep.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(18)
        r = p.add_run(ln)
        r.font.name, r.font.size = "Consolas", Pt(9)


def _quote(rep: Report, text: str) -> None:
    rep.doc.add_paragraph(text, style="Quote")


def _plan_why(run: Path) -> dict[str, str]:
    f = run / "plan.md"
    out = {}
    if f.is_file():
        for ln in f.read_text(encoding="utf-8").splitlines():
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            if len(cells) >= 4 and cells[0].isdigit():
                out[cells[1]] = cells[3]
    return out


def _plan_notes(run: Path) -> list[str]:
    f = run / "plan.md"
    if not f.is_file():
        return []
    text = f.read_text(encoding="utf-8")
    return [re.sub(r"\*\*", "", ln).strip() for ln in text.splitlines()
            if re.search(r"plan note|why the plan changed|\(v\d|resolved", ln, re.I) and not ln.startswith("#")]


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _prob(x: float | None) -> str:
    """Posterior probabilities from Monte Carlo draws: never print 100.0% / 0.0%."""
    if x is None:
        return "n/a"
    return "> 99.9%" if x > 0.999 else ("< 0.1%" if x < 0.001 else R.fmt_prob(x))


def _evidence(r: dict[str, Any]) -> str:
    if r.get("p_value_adjusted") is not None and r.get("correction") not in (None, "none"):
        return f"{R.fmt_p_stat(r['p_value_adjusted'])} (adj., {r.get('correction')})"
    if r.get("p_value") is not None:
        return R.fmt_p_stat(r["p_value"])
    if r.get("probability") is not None:
        return f"P(beat) {_prob(r['probability'])}"
    return "n/a"


def _prior(p: dict[str, Any] | None) -> str:
    if not p:
        return "n/a"
    if p.get("dist") == "beta":
        return f"Beta({R.fmt_num(p['a'], 1)}, {R.fmt_num(p['b'], 1)})"
    if p.get("dist") == "normal":
        return f"Normal(0, sd {R.fmt_num(p.get('sd'))}) on the difference"
    return str(p.get("dist", "n/a"))


def _clip(text: str, n: int = 240) -> str:
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[:n].rsplit(" ", 1)[0] + " ..."


def _num(x: Any) -> str:
    return R.fmt_num(x) if isinstance(x, (int, float)) else "n/a"


def _estimate_rows(rows: list[dict[str, Any]]) -> list[list[str]]:
    out = []
    for r in rows:
        out.append([r.get("step_id", ""), str(r.get("metric") or ""), str(r.get("segment") or ""),
                    (r.get("method") or "").split(".")[-1].replace("_", " "),
                    f"{_num(r.get('estimate'))} [{_num(r.get('ci_low'))}, {_num(r.get('ci_high'))}]",
                    f"{R.fmt_pct(r.get('rel_lift'))} {R.fmt_ci(r.get('rel_ci_low'), r.get('rel_ci_high'))}" if r.get("rel_lift") is not None else "n/a",
                    _evidence(r), " / ".join(f"{k}: {R.fmt_num(v)}" for k, v in (r.get("n") or {}).items())])
    return out


EFFECT_METHODS = ("ztest", "ttest", "bootstrap", "noninferiority", "beta_binomial", "normal_model", "cuped", "regression",
                  "delta_method", "cluster_robust", "diff_in_diff", "msprt", "custom.", "mann_whitney", "synthetic")
DIAGNOSTIC_METHODS = ("interaction_test", "simpsons_check", "trend_test", "parallel_trends", "srm_check", "chi_square")


# --------------------------------------------------------------------------- builder


def build_technical_report(run_dir: str | Path, out: str | Path | None = None) -> Path:
    """Write ``technical_report.docx`` for a run (see module docstring for the structure)."""
    ctx = run_context(run_dir)
    run, st, res = ctx["run"], ctx["state"], ctx["results"]
    summ = res.get("summary") or {}
    intake = intake_sections(run)
    s = st["settings"]
    allr = R.all_results(res)
    rep = Report()
    rep.title(f"Technical report: {ctx['title']}", "Companion to summary.docx: how the result was obtained, for data scientists",
              ctx["source"])
    rep.footer(f"{ctx['title']} | {run.name} | technical report | numbers from results.json")

    # 1 Technical summary --------------------------------------------------------------------------------------------
    rep.heading("1. Technical summary")
    if summ:
        rep.para(f"{summ.get('verdict_label', 'n/a')}. {summ.get('reason', '')}", bold_lead="Verdict: ")
        rep.para(summ.get("headline", ""), bold_lead="Headline: ")
        rep.para("ship if the primary metric is significantly better and the relative lift is at least the practical "
                 "threshold with no guardrail breach; don't ship if it is significantly worse or a guardrail is breached "
                 "(iterate if the primary also won); iterate if significant but below the threshold, or if not significant "
                 "with the CI excluding the threshold; otherwise extend (inconclusive). Bayesian results use P(beat) against "
                 f"{s.get('bayes_decision_threshold')}.", bold_lead="Decision rule (abkit.decision.recommend): ")
    prim = [r for r in allr if r.get("role") == "primary"]
    if prim:
        rep.table(["Step", "Metric", "Segment", "Method", "Estimate [CI]", "Relative lift [CI]", "Evidence", "n"],
                  _estimate_rows(prim), "Primary estimate", 8)
    if summ.get("caveats"):
        rep.heading("Caveats", 2)
        rep.bullets(summ["caveats"])

    # 2 Problem framing ----------------------------------------------------------------------------------------------
    rep.heading("2. Problem framing and design")
    question = next((b for h, b in intake.items() if h.startswith("business question")), "")
    if question:
        rep.para(" ".join(question.split()), bold_lead="Business question: ")
    pm = next((m for m in st.get("metrics", []) if m.get("role") == "primary"), {})
    if pm:
        par = "p" if pm.get("type") == "binary" else "mu"
        alt = {"two-sided": "!=", "one-sided": ">" if pm.get("direction") in ("increase", "no_decrease") else "<"}.get(s["sidedness"], "!=")
        rep.para("for the primary metric:", bold_lead="Formal hypotheses ")
        _code(rep, [f"H0: {par}_T = {par}_C      H1: {par}_T {alt} {par}_C      ({s['sidedness']}, alpha = {s['alpha']})"])
    design = st.get("design") or {}
    rows = [["Problem types", ", ".join(st.get("problem_type", []))]]
    if design:
        rows += [["Variant column / control / treatment", f"{design.get('variant_col')} / {design.get('control')} / {design.get('treatment')}"],
                 ["Randomisation unit", str(design.get("unit_col"))],
                 ["Intended split", ", ".join(f"{k}: {v}" for k, v in (design.get("expected_split") or {}).items())]]
    rows.append(["Period", ctx["period"] or "not recorded"])
    rep.table(["Item", "Value"], rows, "Design", 9, [1.3, 3])
    mrows = [[m.get("name"), m.get("label", ""), m.get("role"), m.get("type"), m.get("direction"),
              str(m.get("practical_threshold_rel", m.get("ni_margin_rel", "")))] for m in st.get("metrics", [])]
    rep.table(["Metric", "Label", "Role", "Type", "Direction", "Threshold / margin (rel.)"], mrows, "Metrics", 9)
    rep.table(["Setting", "Value"], [[k, str(v)] for k, v in s.items()], "Statistical settings used", 8.5, [2, 2])
    if st.get("overrides"):
        rep.bullets([f"{k}: {v['from']} -> {v['to']} ({v.get('at', '')})" for k, v in st["overrides"].items()])

    # 3 Data and validation ------------------------------------------------------------------------------------------
    rep.heading("3. Data and validation")
    lineage = []
    if st.get("source_csv"):
        lineage.append(["Original file", str(st["source_csv"])])
    if st.get("input_csv") and (run / st["input_csv"]).is_file():
        lineage += [["Run copy (analysed)", st["input_csv"]], ["SHA-256 of run copy", _sha256(run / st["input_csv"])]]
    schema = next((c for c in (res.get("validation") or {}).get("checks", []) if c["check"] == "schema"), None)
    if schema:
        lineage.append(["Shape", str(schema["detail"])])
    if lineage:
        rep.table(["Item", "Value"], lineage, "Data lineage (the original file is never modified)", 8.5, [1.2, 3.5])
    if res.get("validation"):
        rep.para(f"overall {res['validation'].get('verdict', 'n/a').upper()} (script {res['validation'].get('script', 'n/a')}). "
                 "Outliers are reported, never removed automatically.", bold_lead="Validation verdict: ")
        rep.table(["Check", "Verdict", "Detail"], _validation_rows(res), "Validation checks", 8, [1.2, 0.8, 5], 1)
    _figures(rep, ctx, ["validation"])

    # 4 Methods ------------------------------------------------------------------------------------------------------
    rep.heading("4. Methods, step by step")
    why = _plan_why(run)
    for step in st["plan"]["steps"]:
        sid = step["id"]
        rep.heading(f"{sid}: {step['method']} ({step.get('status', 'n/a')})", 2)
        if why.get(sid):
            rep.para(why[sid], bold_lead="Why (plan): ")
        st_res = res["steps"].get(sid) or (res.get("validation") if sid == "validate" else None) or {}
        if st_res.get("script"):
            rep.para(st_res["script"], bold_lead="Script: ")
        methods = list(dict.fromkeys([step["method"]] + [r.get("method", "") for r in (st_res.get("results") or [])]))
        documented = False
        for meth in methods:
            doc = method_doc(meth)
            if not doc:
                continue
            documented = True
            rep.para(f"{doc['what']}", bold_lead=f"{doc['name']}: ")
            if doc["formula"]:
                _code(rep, doc["formula"])
            rep.para(f"{doc['inference']} Assumptions: {doc['assumptions']} Reference: {doc['reference']}.")
        script = run / (st_res.get("script") or "")
        if not documented and script.is_file():
            text = script.read_text(encoding="utf-8", errors="replace")
            m = re.search(r'^\s*(?:"""|\'\'\')(.+?)(?:"""|\'\'\')', text, re.S)
            rep.para(" ".join((m.group(1) if m else "See the script.").split()),
                     bold_lead="Custom method (not in abkit; documented by its script): ")
        rlist = st_res.get("results") or []
        if rlist:
            alphas = {r.get("alpha") for r in rlist if r.get("alpha") is not None}
            corr = {r.get("correction") for r in rlist if r.get("correction")}
            rep.para(f"alpha {', '.join(str(round(a, 4)) for a in alphas) or 'n/a'}; sidedness "
                     f"{', '.join({str(r.get('sidedness')) for r in rlist if r.get('sidedness')}) or 'n/a'}; correction "
                     f"{', '.join(corr) or 'none'}.", bold_lead="Settings applied: ")
            seen, arows = set(), []
            for r in rlist:
                for a in r.get("assumptions", []):
                    key = (a.get("check"), a.get("passed"))
                    if key not in seen:
                        seen.add(key)
                        arows.append([a.get("check", ""), {True: "PASS", False: "FAIL", None: "NOT CHECKED"}[a.get("passed")],
                                      str(a.get("detail", ""))[:120]])
            if arows:
                rep.table(["Assumption", "Status", "Detail"], arows, f"Assumptions checked in step {sid}", 8, [2.2, 0.8, 3])
            notes = list(dict.fromkeys(n for r in rlist for n in r.get("notes", [])))[:6]
            if notes:
                rep.bullets(notes)

    # 5 Results ------------------------------------------------------------------------------------------------------
    rep.heading("5. Results")
    eff = [r for r in allr if any(k in (r.get("method") or "") for k in EFFECT_METHODS)]
    for i in range(0, len(eff), 25):
        rep.table(["Step", "Metric", "Segment", "Method", "Estimate [CI] (abs.)", "Relative [CI]", "Evidence", "n"],
                  _estimate_rows(eff[i:i + 25]), "All effect estimates" + (" (cont.)" if i else ""), 7)
    bayes = [r for r in allr if (r.get("method") or "").startswith("bayesian.")]
    if bayes:
        rep.table(["Metric", "P(beat)", "Expected loss (treatment)", "Expected loss (control)", "Credible level", "Prior"],
                  [[str(r.get("metric")), _prob(r.get("probability")),
                    _num((r.get("extra") or {}).get("expected_loss_treatment")), _num((r.get("extra") or {}).get("expected_loss_control")),
                    R.fmt_prob((r.get("extra") or {}).get("credible_level")), _prior((r.get("extra") or {}).get("prior"))]
                   for r in bayes], "Bayesian outputs", 8)
    diag = [r for r in allr if any(k in (r.get("method") or "") for k in DIAGNOSTIC_METHODS)]
    if diag:
        rep.table(["Step", "Test", "Statistic / estimate", "Evidence", "Interpretation"],
                  [[r.get("step_id", ""), (r.get("method") or "").split(".")[-1].replace("_", " "), _num(r.get("estimate")),
                    _evidence(r), _clip(" ".join(r.get("notes", [])))] for r in diag], "Diagnostic tests", 8)
    _figures(rep, ctx, ["results", "secondary", "guardrails", "segments", "time", "power"])

    # 6 Robustness ---------------------------------------------------------------------------------------------------
    rep.heading("6. Robustness and sensitivity")
    by_metric: dict[str, list[dict[str, Any]]] = {}
    for r in eff:
        if r.get("segment") in (None, "", "All") and r.get("role") != "segment":
            by_metric.setdefault(str(r.get("metric")), []).append(r)
    multi = {k: v for k, v in by_metric.items() if len(v) > 1}
    if multi:
        for metric, rows_ in multi.items():
            rep.table(["Step", "Metric", "Segment", "Method", "Estimate [CI] (abs.)", "Relative [CI]", "Evidence", "n"],
                      _estimate_rows(rows_), f"{metric}: the same metric across methods", 7)
    else:
        rep.para("Only one method was applied per metric; no cross-method comparison is available.")
    if prim and summ.get("mde_rel") is not None:
        p = prim[0]
        thr = summ["mde_rel"]
        lines = [f"Relative lift {R.fmt_pct(p.get('rel_lift'))}; CI lower bound {R.fmt_pct(p.get('rel_ci_low'))} vs the "
                 f"{R.fmt_pct(thr, 0, signed=False)} practical threshold: the whole CI is "
                 + ("above it." if (p.get("rel_ci_low") or 0) >= thr else "not above it (the true lift may be below the threshold).")]
        if p.get("p_value") is not None:
            lines.append(f"Evidence {R.fmt_p_stat(p['p_value'])} vs alpha {s['alpha']}.")
        rep.heading("Distance to the decision boundaries", 2)
        rep.bullets(lines)
    rep.heading("What could change the conclusion", 2)
    rep.bullets(summ.get("caveats") or ["No caveats recorded."])

    # 7 Decision log -------------------------------------------------------------------------------------------------
    rep.heading("7. Decision log")
    rep.para("Verbatim records (in quote style) may mention values that were later superseded; the current numbers are "
             "those in sections 1-6.")
    qa = next((b for h, b in intake.items() if h.startswith("q&a")), "")
    if qa:
        rep.heading("Intake questions and answers", 2)
        for ln in qa.splitlines():
            if ln.strip():
                _quote(rep, re.sub(r"\*\*", "", ln.strip()))
    decisions = [(h, b) for h, b in intake.items() if h.startswith("decision ")]
    if decisions:
        rep.heading("User decisions at gates", 2)
        for h, b in decisions:
            _quote(rep, f"{h.capitalize()}: " + " ".join(re.sub(r"\*\*", "", b).split()))
    notes = _plan_notes(run)
    if notes:
        rep.heading("Plan changes", 2)
        for n in notes:
            _quote(rep, n)
    if st.get("history"):
        rep.heading("State history", 2)
        rep.table(["When (UTC)", "Event"], [[h.get("at", ""), h.get("event", "")] for h in st["history"]], "state.json history", 8, [1.3, 3])
    review = run / "review.md"
    if review.is_file():
        rep.heading("Review rounds", 2)
        for block in [b for b in re.split(r"^(?=# Review)", review.read_text(encoding="utf-8"), flags=re.M) if b.strip()]:
            lines = block.strip().splitlines()
            _quote(rep, lines[0].lstrip("# ").strip() + " - " + next((ln for ln in lines if ln.lower().startswith("verdict")), ""))
            for ln in lines:
                cells = [c.strip() for c in ln.strip().strip("|").split("|")]
                if ln.strip().startswith("|") and len(cells) >= 3 and cells[0].isdigit():
                    _quote(rep, f"#{cells[0]} [{cells[1]}] {cells[2]}" + (f" | fix: {cells[4]}" if len(cells) > 4 and cells[4] else ""))

    # 8 Reproducibility ----------------------------------------------------------------------------------------------
    rep.heading("8. Reproducibility")
    pkgs = []
    for name in ("pandas", "numpy", "scipy", "statsmodels", "matplotlib", "python-pptx", "python-docx", "PyYAML"):
        try:
            pkgs.append(f"{name} {metadata.version(name)}")
        except metadata.PackageNotFoundError:
            pkgs.append(f"{name} (not installed)")
    rep.table(["Item", "Value"], [["Python", sys.version.split()[0]], ["Platform", platform.platform()],
                                  ["Packages", ", ".join(pkgs)], ["abkit", __import__("abkit").__version__],
                                  ["results.json schema", str(res.get("schema_version"))],
                                  ["Seeds", f"bootstrap seed {s.get('bootstrap_seed')} ({s.get('bootstrap_n_resamples')} resamples); "
                                            "Bayesian draws seed 42 (abkit default)"]],
              "Environment", 8.5, [1.2, 3.5])
    scripts = sorted((run / "code").glob("*.py"))
    produced = {}
    for sid, stp in res.get("steps", {}).items():
        produced.setdefault(Path(stp.get("script", "")).name, []).append(sid)
    if res.get("validation", {}).get("script"):
        produced.setdefault(Path(res["validation"]["script"]).name, []).append("validation")
    if summ.get("script"):
        produced.setdefault(Path(summ["script"]).name, []).append("summary")
    rep.table(["Script", "Purpose", "Writes results for"],
              [[p.name, _first_doc_line(p), ", ".join(produced.get(p.name, [])) or "-"] for p in scripts], "Code index", 8, [1.6, 3, 1.2])
    rep.heading("How to rerun", 2)
    rep.para("From the repository root, in this order (the plan_gate hook requires the run's gates to be open for 20_-50_ scripts):")
    _code(rep, [f"python runs/{run.name}/code/{p.name}" for p in scripts])
    rep.para("Then check every printed number against results.json:")
    _code(rep, [f"python -m abkit.tracecheck {run.name}"])
    log = run / "run_log.jsonl"
    if log.is_file():
        entries = [json.loads(ln) for ln in log.read_text(encoding="utf-8").splitlines() if ln.strip()]
        failed = sum(e.get("status") == "failed" for e in entries)
        rep.para(f"{len(entries)} commands logged while this run was active ({failed} failed); see run_log.jsonl.",
                 bold_lead="Run log: ")
    arts = [[str(p.relative_to(run)).replace("\\", "/"), R.fmt_num(p.stat().st_size / 1024, 1) + " KB"]
            for p in sorted(run.rglob("*")) if p.is_file() and p.suffix in (".json", ".md", ".pptx", ".docx", ".csv", ".png")
            and not p.name.startswith("~$")]
    rep.table(["Artefact", "Size"], arts[:60], "Run artefacts", 8, [4, 1])
    cands = res.get("candidates_for_promotion") or []
    if cands:
        rep.heading("Promotion candidates", 2)
        rep.bullets([f"{c['name']} ({c['file']}) -> {c.get('target_module') or '?'}: {c['description']}" for c in cands])
    _figures(rep, ctx, ["appendix"])

    return rep.save(Path(out) if out else run / "technical_report.docx")
