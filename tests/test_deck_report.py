"""Smoke tests: build a deck and a report from a sample run and check structure and numbers."""

import json
import re
import sys

import pytest
from docx import Document
from pptx import Presentation

from abkit import results as R
from abkit.deck import Deck, build_standard_deck
from abkit.report import build_standard_report


@pytest.fixture(scope="module")
def sample_run(tmp_path_factory, demo_csvs):
    import shutil

    from conftest import ROOT

    root = tmp_path_factory.mktemp("proj")
    shutil.copytree(ROOT / "config", root / "config")
    (root / "runs").mkdir()
    sys.path.insert(0, str(demo_csvs))
    from sample_run import build_checkout_sample

    return build_checkout_sample(root=root)


def deck_text(prs):
    return "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame) + "\n" + "\n".join(
        c.text for s in prs.slides for sh in s.shapes if sh.has_table for row in sh.table.rows for c in row.cells)


def test_deck_structure_and_numbers(sample_run):
    path = build_standard_deck(sample_run)
    prs = Presentation(str(path))
    assert prs.slide_width / prs.slide_height == pytest.approx(16 / 9, rel=0.01)
    # title, exec summary, agenda, ask, approach | findings: 3 results charts, guardrail chart + table, segment chart |
    # impact, risks, recommendation | appendix: divider, methodology, validation table, full results
    assert len(prs.slides) == 18
    res = json.loads((sample_run / "results.json").read_text())
    text = deck_text(prs)
    assert res["summary"]["headline"] in text
    for k in res["summary"]["key_numbers"]:
        assert k["value"] in text
    p = res["steps"]["primary_test"]["results"][0]
    assert R.fmt_pct(p["rel_lift"]) in text and R.fmt_ci(p["rel_ci_low"], p["rel_ci_high"]) in text
    for s in list(prs.slides)[1:]:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if "Appendix" in texts:          # section divider
            continue
        assert s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()


def test_storyline_is_fixed_and_answer_first(sample_run):
    from abkit.deck.builder import CHAPTERS

    build_standard_deck(sample_run)
    story = (sample_run / "deck_storyline.md").read_text(encoding="utf-8")
    rows = [r.split(" | ") for r in story.splitlines() if r.startswith("| ") and not r.startswith("| #")]
    chapters = [r[1].strip() for r in rows]
    titles = [r[2].rstrip(" |") for r in rows]
    res = R.load_results(sample_run)
    assert titles[1] == res["summary"]["headline"]            # executive summary (answer first) is slide 2
    assert titles[2] == "Agenda"
    seen = [c for i, c in enumerate(chapters) if c in CHAPTERS and (i == 0 or chapters[i - 1] != c)]
    assert seen == CHAPTERS                                     # the five chapters, in order, each once
    assert titles[3].startswith("The ask")
    assert not (sample_run / "deck_warnings.txt").exists()


def test_next_step_parsing_and_impact_fallback(sample_run):
    from abkit.deck.builder import _impact_content, _parse_step, run_context

    assert _parse_step("Data owner: confirm random assignment (within 1 week)") == (
        "Confirm random assignment", "Data owner", "within 1 week")
    assert _parse_step("Roll out to everyone") == ("Roll out to everyone", "To agree", "To agree")
    assert _parse_step({"action": "ship it", "owner": "PM"}) == ("Ship it", "PM", "To agree")
    title, tiles, so_what, foot = _impact_content(run_context(sample_run))
    assert len(tiles) == 3 and tiles[0][0] == R.fmt_pct(R.get_result(sample_run, "primary_test")["rel_lift"])
    assert so_what


def test_deck_numbers_all_trace_to_results(sample_run):
    """Every percentage printed in the deck appears among values formatted from results.json."""
    prs = Presentation(str(build_standard_deck(sample_run)))
    res = R.load_results(sample_run)
    allowed = set()
    for r in R.all_results(res) + [res["steps"]["segments"]["data"]["overall"]]:
        for k in ("rel_lift", "rel_ci_low", "rel_ci_high", "control_value", "treatment_value"):
            if r.get(k) is not None:
                allowed |= {R.fmt_pct(r[k]), R.fmt_pct(r[k], signed=False, decimals=2), R.fmt_pct(r[k], 0),
                            R.fmt_pct(abs(r[k]), signed=False)}
        if r.get("alpha"):
            allowed.add(f"{round((1 - r['alpha']) * 100)}%")
        m = (r.get("extra") or {}).get("margin_rel")
        if m is not None:
            allowed |= {f"{m:.0%}", f"{m:.1%}"}
    for c in res["validation"]["checks"]:
        allowed |= set(re.findall(r"[\d.]+%", str(c["detail"])))
    printed = set(re.findall(r"[+\u2212-]?\d+(?:\.\d+)?%", deck_text(prs)))
    allowed |= {f"{res['summary']['mde_rel']:.0%}", "100%"}          # practical threshold; rollout wording in next steps
    unknown = {x for x in printed if x not in allowed and x.lstrip("+") not in allowed}
    assert not unknown, f"numbers in deck not traceable to results.json: {unknown}"


def test_report_sections_and_captions(sample_run):
    path = build_standard_report(sample_run)
    doc = Document(str(path))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1 == ["Executive summary", "Business question and hypothesis", "Experiment design", "Data and validation",
                  "Methodology", "Results", "Guardrails", "Segments", "Risks, limitations and assumptions",
                  "Recommendation and next steps", "Appendix"]
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert sum(c.startswith("Figure") for c in captions) == 5
    assert sum(c.startswith("Table") for c in captions) == len(doc.tables) - 1       # verdict box is a 1-cell table
    body = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    res = R.load_results(sample_run)
    assert res["summary"]["headline"] in body
    for k in res["summary"]["key_numbers"]:
        assert k["value"] in body
    assert doc.styles["Normal"].font.name == "Arial"


def test_deck_sizes_with_plan():
    """A power-analysis style deck with few sections stays short; word-count warnings fire on long text."""
    d = Deck("src")
    d.title_slide("Power analysis", "How long to run")
    d.bullets("Too many words", ["word " * 80], "notes")
    assert len(d.prs.slides) == 2 and d.warnings


def test_missing_summary_raises(sample_run):
    res = R.load_results(sample_run)
    saved = res["summary"]
    res["summary"] = {}
    R.save_results(sample_run, res)
    try:
        with pytest.raises(ValueError):
            build_standard_deck(sample_run)
    finally:
        res["summary"] = saved
        R.save_results(sample_run, res)


def test_tracecheck_passes_on_sample_run(sample_run):
    from abkit import tracecheck

    build_standard_deck(sample_run)
    build_standard_report(sample_run)
    problems = tracecheck.check(sample_run)
    assert not [p for p in problems if "100%" not in p], problems
